"""Durable PostgreSQL executor. Never imports the legacy HTTP server or SQLite jobs."""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
import signal
import uuid
from datetime import date, datetime, timezone
from dataclasses import replace

import httpx
import psycopg
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from openpyxl import Workbook

from app.config import Settings
from app.services.openresults.metadata import EventMetadataService
from app.services.openresults.catalog import EventCatalog
from app.services.scraper import OpenResultsScraper
from batch import BatchRun
from presence import announce, stop_requested


def storage_headers():
    key=os.environ.get('SUPABASE_SECRET_KEY') or os.environ['SUPABASE_SERVICE_ROLE_KEY']
    return {'apikey':key, **({} if key.startswith('sb_secret_') else {'Authorization':f'Bearer {key}'})}


def claim_next_task():
    file=os.environ.get('WORKER_TASK_SELECTION_FILE')
    if not file:
        return query('SELECT * FROM claim_task(%s,%s)',(['openresults','exports'],str(uuid.uuid4())),True)
    try:
        from pathlib import Path
        ids=json.loads(Path(file).read_text(encoding='utf-8'))
        if not isinstance(ids,list) or len(ids)>100 or any(not isinstance(i,str) or not i or len(i)>100 for i in ids):raise ValueError()
    except Exception:
        raise ValueError('task_selection_invalid') from None
    return query('SELECT * FROM claim_selected_task(%s,%s,%s)',(['openresults','exports'],str(uuid.uuid4()),ids),True)


def connection():
    # psycopg uses libpq URI, not Prisma's ?schema=public parameter.
    return psycopg.connect(os.environ["WORKER_DATABASE_URL"], row_factory=dict_row)


def query(sql, params=(), one=False):
    with connection() as conn:
        cur = conn.execute(sql, params)
        return cur.fetchone() if one else cur.fetchall()


def fenced(conn, task):
    row = conn.execute('''SELECT id FROM "CollectionTask" WHERE id=%s AND status='running'
        AND "leaseToken"=%s AND "leaseUntil">now() FOR UPDATE''', (task['id'], task['leaseToken'])).fetchone()
    if not row:
        raise RuntimeError('lease_lost')


def publish(task, result):
    """Publish a complete replacement atomically; a failed transaction leaves old rows intact."""
    if result.warnings or not result.records or (result.expected_total is not None and len(result.records) != result.expected_total):
        raise ValueError('incomplete_extraction')
    payload = task['payload']
    records = result.records
    canonical = json.dumps(records, sort_keys=True, default=str, ensure_ascii=False)
    with connection() as conn:
        fenced(conn, task)
        reference = conn.execute('''SELECT r."eventId",e.date FROM "EventSourceReference" r JOIN "Event" e ON e.id=r."eventId"
            WHERE r."sourceType"='openresults' AND r."sourceExternalId"=%s FOR UPDATE OF r''', (payload['externalId'],)).fetchone()
        if not reference or reference['eventId'] != payload['eventId']:
            raise ValueError('association_changed')
        if not result.metadata.event_date or not reference['date'] or reference['date'].date() != result.metadata.event_date:
            raise ValueError('edition_date_mismatch')
        if result.metadata.event_id and str(result.metadata.event_id) != payload['externalId']:
            raise ValueError('source_identity_mismatch')
        result_set = conn.execute('''INSERT INTO "ResultSet" (id,"eventId",source,"externalId","sourceUrl","updatedAt","contentHash",count)
            VALUES (%s,%s,'openresults',%s,%s,now(),%s,%s) ON CONFLICT (source,"externalId") DO UPDATE
            SET "updatedAt"=now(),"contentHash"=EXCLUDED."contentHash",count=EXCLUDED.count,"sourceUrl"=EXCLUDED."sourceUrl" RETURNING id''',
            (str(uuid.uuid4()),payload['eventId'],payload['externalId'],payload['url'],hashlib.sha256(canonical.encode()).hexdigest(),len(records))).fetchone()['id']
        conn.execute('DELETE FROM "RaceResult" WHERE "resultSetId"=%s',(result_set,))
        conn.execute('DELETE FROM "RaceDiscipline" WHERE "resultSetId"=%s',(result_set,))
        for modality in result.modalities:
            conn.execute('''INSERT INTO "RaceDiscipline" (id,"resultSetId","externalId",name) VALUES (%s,%s,%s,%s)
                ON CONFLICT ("resultSetId","externalId") DO NOTHING''',(str(uuid.uuid5(uuid.NAMESPACE_URL,result_set+":"+str(modality.value))),result_set,str(modality.value),modality.name))
        for record in records:
            # Scoped to one source/edition; never creates a person identity from a name.
            digest=hashlib.sha256(json.dumps({k:record.get(k) for k in ['modality','gender','bib','name','category','overall_position']},sort_keys=True).encode()).hexdigest()
            conn.execute('''INSERT INTO "RaceResult" (id,"resultSetId","recordKey",modality,gender,category,bib,name,team,"overallPosition","categoryPosition",time,pace,"distanceKm",gap)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
                (str(uuid.uuid5(uuid.NAMESPACE_URL,result_set+":"+digest)),result_set,digest,str(record.get('modality','')),record.get('gender'),record.get('category'),str(record.get('bib','')),
                 record.get('name',''),record.get('team'),position(record.get('overall_position')),position(record.get('category_position')),record.get('time'),record.get('pace'),record.get('distance_km'),record.get('gap')))
        conn.execute('UPDATE "EventSourceReference" SET "lastSeenAt"=now(),"updatedAt"=now() WHERE "sourceType"=\'openresults\' AND "sourceExternalId"=%s',(payload['externalId'],))
        conn.execute('SELECT finish_task(%s,%s,\'completed\',%s,NULL)',(task['id'],task['leaseToken'],Jsonb({'processed':len(records),'stage':'published'})))


def position(value):
    try:
        number=int(value)
        return number if number>0 else None
    except (TypeError,ValueError):
        return None


def store_match(task, metadata):
    external_id=str(metadata.event_id) if metadata.event_id else 'url:'+hashlib.sha256(metadata.source_url.encode()).hexdigest()
    with connection() as conn:
        fenced(conn,task)
        conn.execute('''INSERT INTO "SourceMatch" (id,source,"externalId",url,name,date,city,state,status,"updatedAt")
            VALUES (%s,'openresults',%s,%s,%s,%s,%s,%s,'pending',now()) ON CONFLICT (source,"externalId") DO UPDATE SET
            url=EXCLUDED.url,name=EXCLUDED.name,date=EXCLUDED.date,city=EXCLUDED.city,state=EXCLUDED.state,"updatedAt"=now()''',
            (str(uuid.uuid4()),external_id,metadata.source_url,metadata.name,metadata.event_date,metadata.city,metadata.state))
        # Only an existing exact source identity is reused automatically. Names are not sufficient.
        conn.execute('''UPDATE "SourceMatch" m SET status='resolved',"eventId"=r."eventId","resolvedBy"='exact_reference'
            FROM "EventSourceReference" r JOIN "Event" e ON e.id=r."eventId"
            WHERE m.source=r."sourceType" AND m."externalId"=r."sourceExternalId" AND (m.date AT TIME ZONE 'UTC')::date=e.date::date
            AND m.source='openresults' AND m."externalId"=%s''',(external_id,))


def excel_cell(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):
        return "'"+value
    return value


def build_workbook(event_id):
    workbook=Workbook(write_only=True)
    sheet=workbook.create_sheet('Resultados')
    columns=['name','bib','modality','gender','category','overallPosition','categoryPosition','time','pace','team','source','sourceUrl','updatedAt']
    sheet.append(columns)
    index=0
    with connection() as conn:
        with conn.cursor(name='export_rows') as cursor:
            cursor.execute('''SELECT r.*,s.source,s."sourceUrl",s."updatedAt" FROM "RaceResult" r JOIN "ResultSet" s ON s.id=r."resultSetId"
                JOIN "Event" e ON e.id=s."eventId" WHERE s."eventId"=%s AND e."publicationStatus"='published' ORDER BY r.id''',(event_id,))
            for index,row in enumerate(cursor,1):
                if index%1000000==0:
                    sheet=workbook.create_sheet(f'Resultados-{index}');sheet.append(columns)
                sheet.append([excel_cell(str(row[c]) if isinstance(row[c],datetime) else row[c]) for c in columns])
    if not index:
        raise ValueError('results_unavailable')
    output=io.BytesIO();workbook.save(output)
    if output.tell()>50*1024*1024:
        raise ValueError('export_too_large')
    return output


async def export(task):
    if task['kind']=='export-selection':
        artifact=query('SELECT * FROM "ExportArtifact" WHERE "taskId"=%s',(task['id'],),one=True)
        if not artifact:raise ValueError('export_artifact_missing')
    else:
        # Repair the short API enqueue/artifact creation gap after an interrupted request.
        artifact=query('''INSERT INTO "ExportArtifact" (id,"eventId","ownerId","taskId",status,"expiresAt","createdAt")
            VALUES (%s,%s,%s,%s,'queued',%s::timestamptz+interval '1 day',%s)
            ON CONFLICT ("taskId") DO UPDATE SET "taskId"=EXCLUDED."taskId" RETURNING *''',
            (str(uuid.uuid4()),task['payload']['eventId'],task['ownerId'],task['id'],task['createdAt'],task['createdAt']),one=True)
    if artifact['expiresAt'].replace(tzinfo=timezone.utc)<=datetime.now(timezone.utc):
        raise ValueError('export_expired')
    from selection_export import build_selection
    if task['kind']=='export':
        artifact['selection']={'eventIds':[artifact['eventId']],'layout':'consolidated','administrative':False}
        artifact['kind']='results'
    output,extension,content_type,count=await asyncio.to_thread(build_selection,artifact,connection)
    with connection() as conn:
        fenced(conn,task)
    if artifact['expiresAt']<=datetime.now(timezone.utc):
        raise ValueError('export_expired')
    path=f"{artifact['id']}/{task['leaseToken']}.{extension}"
    base=os.environ['SUPABASE_URL'].rstrip('/')
    async with httpx.AsyncClient(timeout=30) as client:
        response=await client.post(f'{base}/storage/v1/object/race-exports/{path}',headers={
            **storage_headers(),
            'Content-Type':content_type,'x-upsert':'true'},content=output.getvalue())
        response.raise_for_status()
    with connection() as conn:
        fenced(conn,task)
        conn.execute('UPDATE "ExportArtifact" SET status=\'completed\',"objectPath"=%s,"contentType"=%s WHERE id=%s',(path,content_type,artifact['id']))
        conn.execute('SELECT finish_task(%s,%s,\'completed\',%s,NULL)',(task['id'],task['leaseToken'],Jsonb({'stage':'exported','exportId':artifact['id'],'processed':count,'format':extension})))


_cleanup_cursor = ''

async def cleanup_exports():
    global _cleanup_cursor
    # Revisit expired artifacts in bounded keyset batches. This also removes
    # orphan uploads from a worker that lost its lease after uploading.
    items=query('SELECT id,"objectPath" FROM "ExportArtifact" WHERE "expiresAt"<now() AND id>%s ORDER BY id LIMIT 100',(_cleanup_cursor,))
    if not items:
        _cleanup_cursor=''
        return
    base=os.environ['SUPABASE_URL'].rstrip('/')+'/storage/v1/object'
    headers=storage_headers()
    async with httpx.AsyncClient(timeout=30) as client:
        for item in items:
            response=await client.post(base+'/list/race-exports',headers=headers,json={'prefix':item['id']+'/', 'limit':100,'offset':0})
            response.raise_for_status()
            names=[row['name'] for row in response.json() if row.get('name','').endswith(('.xlsx','.zip')) and '/' not in row['name']]
            paths=[item['id']+'/'+name for name in names]
            if item['objectPath'] and item['objectPath'] not in paths: paths.append(item['objectPath'])
            if paths:
                response=await client.request('DELETE',base+'/race-exports',headers=headers,json={'prefixes':paths})
                response.raise_for_status()
            with connection() as conn:
                conn.execute("UPDATE \"ExportArtifact\" SET status='expired',\"objectPath\"=NULL WHERE id=%s",(item['id'],))
            _cleanup_cursor=item['id']


async def execute(task):
    progress={'stage':'starting','percent':0}
    async def report(_stage,percent):
        # Avoid logging event names, URLs or athlete data from upstream messages.
        progress.update(stage='extracting',percent=percent)
    async def heartbeat():
        while True:
            await asyncio.sleep(20)
            result=await asyncio.to_thread(query,'SELECT heartbeat_task(%s,%s,%s) AS ok',(task['id'],task['leaseToken'],Jsonb(progress)),True)
            if not result['ok']:
                raise RuntimeError('lease_lost')
    async def work():
        settings=replace(Settings.from_env(),scrape_concurrency=1,catalog_concurrency=1,metadata_concurrency=1)
        if task['kind']=='catalog-sync':
            from catalog_sync import sync_catalog
            await sync_catalog(task,settings,connection,fenced,progress)
        elif task['kind']=='discover':
            payload=task['payload']
            settings=replace(settings,catalog_max_pages=min(10,int(payload.get('maxPages',1))))
            catalog=await EventCatalog(settings).discover(date_from=date.fromisoformat(payload['from']),date_to=date.fromisoformat(payload['to']) if payload.get('to') else None)
            failed=0;processed=0
            for event in catalog.events[:min(100,int(payload.get('limit',25)))]:
                try:
                    metadata,_=await EventMetadataService(settings).fetch(event.event_url)
                    await asyncio.to_thread(store_match,task,metadata)
                except Exception:
                    failed+=1
                processed+=1
                progress.update(stage='discovering',processed=processed,failed=failed,discovered=len(catalog.events))
                await asyncio.sleep(0.5)
            if failed or catalog.warnings or len(catalog.events)>int(payload.get('limit',25)):
                await asyncio.to_thread(query,'SELECT finish_task(%s,%s,\'partial\',%s,%s)',(task['id'],task['leaseToken'],Jsonb(progress),'discovery_partial'))
        elif task['kind']=='inspect':
            metadata,_=await EventMetadataService(settings).fetch(task['payload']['url'])
            from edition_metadata import update_edition
            await asyncio.to_thread(update_edition,task,metadata,connection,fenced)
            await asyncio.to_thread(store_match,task,metadata)
        elif task['kind']=='extract':
            result=await OpenResultsScraper(settings).scrape(task['payload']['url'],report)
            if str(task['payload'].get('externalId','')).startswith('url:'):
                from edition_metadata import update_edition
                await asyncio.to_thread(update_edition,task,result.metadata,connection,fenced)
            await asyncio.to_thread(publish,task,result)
        elif task['kind'] in ('export','export-selection'):
            await export(task)
        else:
            raise ValueError('unsupported_task')
    pulse=asyncio.create_task(heartbeat());job=asyncio.create_task(work())
    try:
        async with asyncio.timeout(1800):
            done,_=await asyncio.wait([pulse,job],return_when=asyncio.FIRST_COMPLETED)
            for item in done:
                await item
        await asyncio.to_thread(query,'SELECT finish_task(%s,%s,\'completed\',%s,NULL)',(task['id'],task['leaseToken'],Jsonb(progress)))
    except Exception as exc:
        from app.models import AccessBlockedError
        if isinstance(exc, AccessBlockedError):
            with connection() as conn:
                conn.execute('UPDATE "CollectionTask" SET "maxAttempts"=attempt WHERE id=%s AND "leaseToken"=%s',(task['id'],task['leaseToken']))
        code=('source_access_blocked' if isinstance(exc, AccessBlockedError) else
              str(exc) if isinstance(exc,ValueError) and str(exc) in {'incomplete_extraction','catalog_pagination_not_advancing','catalog_end_unconfirmed','selected_edition_without_results','export_too_large_refine_selection','export_expired','source_identity_already_associated','edition_date_mismatch'} else 'collection_failed')
        outcome='partial' if code=='incomplete_extraction' else 'failed'
        await asyncio.to_thread(query,'SELECT finish_task(%s,%s,%s,%s,%s)',(task['id'],task['leaseToken'],outcome,Jsonb(progress),code))
    finally:
        pulse.cancel();job.cancel()
        await asyncio.gather(pulse,job,return_exceptions=True)


async def main():
    run = BatchRun()
    watchdog = run.watchdog()
    stopped=False
    def stop(*_):
        nonlocal stopped
        stopped=True
    signal.signal(signal.SIGTERM,stop);signal.signal(signal.SIGINT,stop)
    tick=0
    worker_id=str(uuid.uuid4())
    async def presence_pulse():
      nonlocal stopped
      while not stopped:
        await asyncio.sleep(20)
        try:
          await asyncio.to_thread(announce,query,worker_id,'stopping' if stop_requested() else 'busy' if run.active_task_id else 'available',run.active_task_id)
        except Exception:
          stopped=True
          print('Results: connection lost; stopping before the next task.',flush=True)
    presence_task=None
    try:
      await asyncio.to_thread(announce,query,worker_id,'available')
      presence_task=asyncio.create_task(presence_pulse())
      print('Results executor connected; waiting for panel requests.',flush=True)
      # Batch runs may finish before the old 60-tick cleanup cadence.
      if run.batch and os.environ.get('SUPABASE_URL'):
        try: await cleanup_exports()
        except (httpx.HTTPError,KeyError): print('Export cleanup pending; verify Storage configuration.',flush=True)
      while not stopped and not stop_requested() and run.can_claim():
        task=await asyncio.to_thread(claim_next_task)
        if task:
            run.claimed += 1
            run.active_task_id = task['id']
            await asyncio.to_thread(announce,query,worker_id,'busy',task['id'])
            print(f"Results task {task['id']}: started.",flush=True)
            await execute(task)
            outcome=await asyncio.to_thread(query,'SELECT status FROM "CollectionTask" WHERE id=%s',(task['id'],),True)
            run.tasks.append({'id':task['id'],'status':outcome['status']})
            run.tasks = run.tasks[-100:]
            run.active_task_id = None
            await asyncio.to_thread(announce,query,worker_id,'available')
            print(f"Results task {task['id']}: {outcome['status']}.",flush=True)
        else:
            if run.batch:
                run.reason = 'queue_empty'
                break
            await asyncio.sleep(1)
        tick+=1
        if tick%60==0:
            try: await cleanup_exports()
            except (httpx.HTTPError,KeyError): print('Export cleanup pending; verify Storage configuration.',flush=True)
    except Exception:
        run.reason = 'worker_failed'
        raise RuntimeError('worker_failed') from None
    finally:
        if presence_task:
            presence_task.cancel()
            await asyncio.gather(presence_task,return_exceptions=True)
        try: await asyncio.to_thread(announce,query,worker_id,'stopped')
        except Exception: pass
        if watchdog: watchdog.cancel()
        run.report()


if __name__=='__main__':
    try:
        asyncio.run(main())
    except Exception:
        print('Worker stopped; inspect sanitized task history and database connectivity.',flush=True)
        raise SystemExit(1) from None
