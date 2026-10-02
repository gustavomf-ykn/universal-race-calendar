"""Controlled PostgreSQL consumer checks after the TypeScript reconciliation fixture."""
import json
import os
import sys
from urllib.parse import urlsplit
from types import SimpleNamespace
from datetime import date

from openpyxl import load_workbook
from app.config import Settings
from app.models import EXPORT_COLUMNS
from event_aliases import payload as canonical_payload
from result_checkpoints import ResultCheckpoints
from selection_export import build_selection
import worker

mode, old_id, canonical_id, task_id = sys.argv[1:]
uri = urlsplit(os.environ['WORKER_DATABASE_URL'])
assert uri.hostname in ('localhost', '127.0.0.1', 'postgres') and uri.path.endswith('_test')
assert old_id.startswith('reconcile-') and canonical_id.startswith('reconcile-')
task = worker.query('SELECT * FROM "CollectionTask" WHERE id=%s', (task_id,), True)
assert task['payload']['eventId'] == old_id
with worker.connection() as conn:
    assert canonical_payload(conn, task['payload'])['eventId'] == canonical_id

if mode == 'read':
    adapter = ResultCheckpoints(task, Settings(endpoint_page_size=100), worker.connection, worker.fenced)
    with worker.connection() as conn:
        assert adapter._root(conn)['eventId'] == canonical_id
    artifact = {'kind': 'results', 'selection': {'eventIds': [old_id, canonical_id], 'administrative': True}}
    content, extension, _, count = build_selection(artifact, worker.connection)
    book = load_workbook(content)
    assert extension == 'xlsx' and count == 1 and book.active.max_row == 2
    assert [cell.value for cell in book.active[1]] == [label for _, label in EXPORT_COLUMNS]
    book.close()
    assert artifact['selection']['eventIds'] == [old_id, canonical_id]
elif mode == 'publish':
    result = SimpleNamespace(warnings=[], expected_total=1, records=[{
        'name': 'Participante sintético atualizado', 'modality': '5k', 'bib': '001',
        'gender': 'F', 'overall_position': 1, 'time': '00:25:00',
    }], modalities=[SimpleNamespace(value='5k', name='5 km')],
        metadata=SimpleNamespace(event_date=date(2040, 10, 10), event_id=task['payload']['externalId']))
    worker.publish(task, result)
    assert task['payload']['eventId'] == old_id
    saved = worker.query('SELECT status,payload FROM "CollectionTask" WHERE id=%s', (task_id,), True)
    assert saved['status'] == 'completed' and saved['payload']['eventId'] == old_id
else:
    raise ValueError('fixture_mode_invalid')
print(json.dumps({'mode': mode, 'alias': 'resolved', 'payload': 'unchanged', 'status': 'passed'}))
