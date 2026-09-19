from zipfile import ZipFile

from openpyxl import load_workbook

from app.services.exporter import export_events_xlsx, export_result_parts, export_xlsx


def test_generates_single_formatted_results_sheet(tmp_path, sample_result) -> None:
    path, filename = export_xlsx(sample_result, tmp_path)
    assert filename == "corrida-teste-2026-resultados.xlsx"
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Resultados"]
    sheet = workbook["Resultados"]
    assert sheet.freeze_panes == "A2"
    assert len(sheet.tables) == 1
    assert sheet.auto_filter.ref is not None
    assert sheet["A1"].font.bold is True
    headers = [cell.value for cell in sheet[1]]
    bib_column = headers.index("Número") + 1
    assert sheet.cell(2, bib_column).value == "007"
    assert sheet.cell(2, bib_column).data_type == "s"
    assert all(cell.value not in {"nan", "None"} for row in sheet.iter_rows() for cell in row)


def test_simple_events_export_has_exact_columns_and_blank_id(tmp_path) -> None:
    rows = [
        {"event_id": None, "name": "Sem resultado", "start_date": "2026-02-01"},
        {"event_id": "123", "name": "Com resultado", "start_date": "2026-01-01"},
    ]
    path, filename = export_events_xlsx(rows, tmp_path, full=False)
    assert filename == "openresults_provas.xlsx"
    sheet = load_workbook(path)["Provas"]
    assert [cell.value for cell in sheet[1]] == ["id", "nome"]
    assert [sheet.cell(2, 1).value, sheet.cell(2, 2).value] == ["123", "Com resultado"]
    assert sheet.cell(3, 1).value is None


def test_result_parts_are_packaged_as_zip(tmp_path, sample_result) -> None:
    first = {**sample_result.records[0], "event_id": "123"}
    second = {**sample_result.records[1], "event_id": "123"}
    path, filename = export_result_parts([[first], [second]], tmp_path)
    assert filename == "openresults_resultados.zip"
    with ZipFile(path) as archive:
        assert archive.namelist() == [
            "openresults_resultados_parte_001.xlsx",
            "openresults_resultados_parte_002.xlsx",
        ]
    assert not list(tmp_path.glob("*.xlsx"))
