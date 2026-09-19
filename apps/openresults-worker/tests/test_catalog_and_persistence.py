from __future__ import annotations

import time
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import Settings
from app.main import create_app
from app.models import EventMetadata, EventSummary, ModalityInfo
from app.services.openresults.catalog import parse_catalog_payload
from app.services.openresults.metadata import parse_related_metadata
from app.storage import SQLiteStorage


CATALOG_HTML = """
<script>totalEventos = 21765;</script>
<a href="/evento/2026-prova-acento/" class="or-event-card-link">
  <article class="or-event-card">
    <div class="or-event-card-date"><span class="day">09</span><span class="month">ago</span><span class="year">2026</span></div>
    <h2 class="or-event-card-title">Prova Acento</h2>
    <div class="or-event-card-meta"><span>Dianópolis - TO</span><span>160 concluintes</span></div>
    <div class="or-event-card-distances"><span class="or-event-distance">5k | 160</span></div>
  </article>
</a>
"""


def test_catalog_parser_reads_real_shape_and_advertised_total() -> None:
    events, total, has_more = parse_catalog_payload(CATALOG_HTML)
    assert total == 21_765
    assert has_more is None
    assert len(events) == 1
    assert events[0].event_date == date(2026, 8, 9)
    assert events[0].city == "Dianópolis"
    assert events[0].state == "TO"
    assert events[0].expected_total == 160
    assert events[0].modalities == [{"name": "5k", "expected_total": 160}]


def test_related_metadata_removes_people_fields() -> None:
    html = """<script type="application/ld+json">{
      "@type":"SportsEvent", "name":"Teste", "participant":[{"name":"Pessoa"}],
      "performer":{"name":"Atleta"}, "location":{"name":"Parque"}
    }</script>"""
    raw = parse_related_metadata(html, "https://roadrunners.run/evento/teste")
    assert raw["name"] == "Teste"
    assert "participant" not in raw
    assert "performer" not in raw
    assert raw["location"]["name"] == "Parque"


@pytest.mark.asyncio
async def test_sqlite_marks_running_checkpoints_as_interrupted(tmp_path) -> None:
    storage = SQLiteStorage(tmp_path / "state.sqlite3")
    await storage.start()
    event = (
        await storage.upsert_events(
            [EventSummary("Teste", date(2026, 1, 1), "Cidade", "SC", "https://openresults.run/evento/teste/", "teste", event_id="123")]
        )
    )[0]
    await storage.create_job("job", "results", {"mode": "single"})
    await storage.add_job_events("job", [storage.event_row_to_dict(await storage.get_event_by_catalog_id(event.catalog_id))])
    await storage.update_job("job", status="running")
    await storage.update_job_event("job", event.catalog_id, status="processing")
    await storage.upsert_group("job", event.catalog_id, "5k", "5k", "F", status="processing", next_offset=100, extracted_total=100)
    await storage.close()

    recovered = SQLiteStorage(tmp_path / "state.sqlite3")
    await recovered.start()
    assert (await recovered.get_job("job"))["status"] == "interrupted"
    assert (await recovered.get_job_events("job"))[0]["status"] == "partial"
    group = await recovered.get_group("job", event.catalog_id, "5k", "F")
    assert group["status"] == "partial"
    assert group["next_offset"] == 100
    await recovered.close()


def test_persistent_results_api_and_export(tmp_path, sample_result) -> None:
    class FakeMetadata:
        async def fetch(self, url, enrich_roadrunners=False):
            metadata = EventMetadata(
                name="CORRIDA TESTE 2026", event_date=date(2026, 7, 25), city="Florianópolis", state="SC",
                source_url=url, slug="corrida-teste-2026", expected_total=2, event_id="123",
            )
            return metadata, [ModalityInfo("5k", "5k", {"F": 1, "M": 1})]

    class FakeScraper:
        def __init__(self, settings):
            pass

        async def scrape(self, url, progress):
            await progress("Teste persistente", 60)
            return sample_result

    settings = Settings(temp_dir=tmp_path, database_path=tmp_path / "api.sqlite3", cleanup_interval_seconds=600)
    app = create_app(settings, scraper_factory=FakeScraper)
    with TestClient(app) as client:
        app.state.jobs.metadata_service = FakeMetadata()
        response = client.post(
            "/api/results/scrape",
            json={"scope": {"mode": "single", "url": "https://openresults.run/evento/corrida-teste-2026/"}},
        )
        assert response.status_code == 202
        job_id = response.json()["job_id"]
        job = None
        for _ in range(100):
            job = client.get(f"/api/jobs/{job_id}").json()
            if job["status"] in {"completed", "completed_with_warnings", "failed"}:
                break
            time.sleep(0.02)
        assert job["status"] == "completed"
        assert job["total_extracted"] == 2
        results = client.get(f"/api/jobs/{job_id}/results", params={"event_id": "123"}).json()
        assert results["filtered_total"] == 2
        assert all(item["event_id"] == "123" for item in results["items"])
        job_events = client.get(f"/api/jobs/{job_id}/events").json()
        assert job_events["items"][0]["status"] == "completed"
        download = client.get(f"/api/jobs/{job_id}/download")
        assert download.status_code == 200
        path = tmp_path / "download.xlsx"
        path.write_bytes(download.content)
        workbook = load_workbook(path)
        assert workbook.sheetnames == ["Resultados"]
        assert workbook["Resultados"]["A1"].value == "event_id"


def test_all_results_requires_explicit_confirmation(tmp_path) -> None:
    app = create_app(Settings(temp_dir=tmp_path, database_path=tmp_path / "api.sqlite3", cleanup_interval_seconds=600))
    with TestClient(app) as client:
        response = client.post("/api/results/scrape", json={"scope": {"mode": "all"}, "confirm_all": False})
    assert response.status_code == 422
