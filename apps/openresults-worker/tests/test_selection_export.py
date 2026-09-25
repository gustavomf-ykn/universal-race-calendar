import io
import zipfile
from contextlib import contextmanager
from datetime import date,datetime,timezone
from openpyxl import load_workbook
import pytest
import selection_export as exports
from app.models import EXPORT_COLUMNS

class Cursor:
    def __init__(self,rows):self.rows=rows
    def __enter__(self):return self
    def __exit__(self,*args):pass
    def execute(self,*args):pass
    def __iter__(self):return iter(self.rows)
    def fetchall(self):return self.rows
class Database:
    def __init__(self,events,results):self.events=events;self.results=results
    def execute(self,sql,args):return Cursor(self.events if 'FROM "Event" WHERE' in sql else [])
    def cursor(self,**kwargs):return Cursor(self.results)

def fixtures():
    events=[dict(id='edition-1',name='São João',date=date(2026,1,1),city='Camboriú',state='SC',publicationStatus='published',sourceExternalId='123',sourceType='openresults',sourceUrl='https://openresults.run/evento/teste/',description=None,locationName=None,address=None,modality='unknown',registrationUrl=None,officialUrl=None,mainImageUrl=None,updatedAt=datetime.now(timezone.utc))]
    rows=[dict(externalId='123',sourceUrl=events[0]['sourceUrl'],updatedAt=events[0]['updatedAt'],modality='5 km',distanceKm=5,gender='F',overallPosition=1,categoryPosition=1,category='Livre',bib='007',name='=HYPERLINK("bad")',team=None,pace='05:00',time='00:25:00',gap=None)]
    @contextmanager
    def connection():yield Database(events,rows)
    return events,rows,connection

def test_portuguese_reference_headers_text_ids_dates_accents_and_formula_protection():
    _,_,connection=fixtures()
    content,extension,mime,count=exports.build_selection({'kind':'results','selection':{'eventIds':['edition-1']}},connection)
    sheet=load_workbook(content).active
    assert [cell.value for cell in sheet[1]]==[label for _,label in EXPORT_COLUMNS]
    assert sheet['B2'].value=='São João' and sheet['D2'].value=='Camboriú'
    assert sheet['L2'].value=='007' and sheet['L2'].number_format=='@'
    assert sheet['M2'].data_type=='s' and sheet['M2'].value.startswith("'")
    assert sheet.freeze_panes=='A2' and sheet['C2'].number_format=='dd/mm/yyyy'
    assert extension=='xlsx' and count==1

def test_individual_editions_are_zipped_and_missing_results_are_explicit():
    events,rows,connection=fixtures();events.append({**events[0],'id':'edition-2'})
    content,extension,_,count=exports.build_selection({'kind':'results','selection':{'eventIds':['edition-1','edition-2'],'layout':'individual'}},connection)
    assert extension=='zip' and count==2
    with zipfile.ZipFile(content) as archive:
        assert len(archive.namelist())==2
        for name in archive.namelist():assert load_workbook(io.BytesIO(archive.read(name))).active.max_row==2
    rows.clear()
    with pytest.raises(ValueError,match='selected_edition_without_results'):
        exports.build_selection({'kind':'results','selection':{'eventIds':['edition-1','edition-2']}},connection)

def test_full_catalog_retains_nulls_and_distinguishes_ids():
    _,_,connection=fixtures()
    content,_,_,count=exports.build_selection({'kind':'catalog-full','selection':{'eventIds':['edition-1'],'administrative':True}},connection)
    sheet=load_workbook(content).active
    assert sheet['A2'].value=='edition-1' and sheet['B2'].value=='123'
    assert sheet['J2'].value is None and count==1
