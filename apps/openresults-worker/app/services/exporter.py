from __future__ import annotations

import re
import unicodedata
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from app.models import EXPORT_COLUMNS, ExportError, ExtractionResult


FULL_EVENT_COLUMNS: list[tuple[str, str]] = [
    ("event_id", "id"),
    ("name", "nome"),
    ("start_date", "data_inicio"),
    ("end_date", "data_fim"),
    ("city", "cidade"),
    ("state", "uf"),
    ("country", "pais"),
    ("expected_total", "total_esperado"),
    ("event_url", "url_openresults"),
    ("about_url", "url_sobre"),
    ("official_result_url", "url_resultado_oficial"),
    ("registration_url", "url_inscricao"),
    ("image_url", "imagem"),
    ("location_name", "local"),
    ("address", "endereco"),
    ("latitude", "latitude"),
    ("longitude", "longitude"),
    ("description", "descricao"),
    ("source_event_status", "status_origem"),
    ("derived_event_status", "status_derivado"),
    ("modalities", "modalidades"),
    ("external_ids", "ids_externos"),
    ("metadata_fetched_at", "metadados_atualizados_em"),
]


def _stringify_nested(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        import json

        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def _styled_workbook(
    rows: list[dict[str, Any]],
    columns: list[tuple[str, str]],
    *,
    sheet_name: str,
    table_name: str,
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    headers = [label for _, label in columns]
    sheet.append(headers)
    for row in rows:
        values: list[Any] = []
        for field, _ in columns:
            value = _stringify_nested(_excel_value(row.get(field)))
            if field in {"event_date", "start_date", "end_date"} and isinstance(value, str) and value:
                try:
                    value = date.fromisoformat(value[:10])
                except ValueError:
                    pass
            elif field in {"extracted_at", "metadata_fetched_at"} and isinstance(value, str) and value:
                try:
                    value = datetime.fromisoformat(value).replace(tzinfo=None)
                except ValueError:
                    pass
            if field in {"bib", "event_id"} and value is not None:
                value = str(value)
            values.append(value)
        sheet.append(values)
    if not rows:
        sheet.append([None] * len(columns))
    last_column = get_column_letter(len(columns))
    reference = f"A1:{last_column}{sheet.max_row}"
    table = Table(displayName=table_name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.auto_filter.ref = reference
    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for index, (field, _) in enumerate(columns, 1):
        if field in {"bib", "event_id"}:
            for row_number in range(2, sheet.max_row + 1):
                sheet.cell(row_number, index).number_format = "@"
        elif field in {"event_date", "start_date", "end_date"}:
            for row_number in range(2, sheet.max_row + 1):
                sheet.cell(row_number, index).number_format = "dd/mm/yyyy"
        elif field in {"extracted_at", "metadata_fetched_at"}:
            for row_number in range(2, sheet.max_row + 1):
                sheet.cell(row_number, index).number_format = "dd/mm/yyyy hh:mm:ss"
    for index, header in enumerate(headers, 1):
        width = len(header)
        for row_number in range(2, min(sheet.max_row, 500) + 1):
            value = sheet.cell(row_number, index).value
            width = max(width, len(str(value)) if value is not None else 0)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 50)
    return workbook


def export_events_xlsx(
    rows: list[dict[str, Any]],
    output_dir: Path,
    *,
    full: bool,
) -> tuple[Path, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    filename = "openresults_provas_completas.xlsx" if full else "openresults_provas.xlsx"
    columns = FULL_EVENT_COLUMNS if full else [("event_id", "id"), ("name", "nome")]
    ordered = sorted(rows, key=lambda item: (item.get("start_date") or "9999-12-31", item.get("name") or ""))
    workbook = _styled_workbook(ordered, columns, sheet_name="Provas", table_name="TabelaProvas")
    path = output_dir / filename
    workbook.save(path)
    return path, filename


def export_result_parts(
    parts: list[list[dict[str, Any]]],
    output_dir: Path,
) -> tuple[Path, str]:
    if not parts or not any(parts):
        raise ExportError("Não há resultados para gerar o arquivo Excel.")
    output_dir.mkdir(parents=True, exist_ok=True)
    core_fields = [field for field, _ in EXPORT_COLUMNS]
    extra_fields = sorted(
        {
            field
            for part in parts
            for row in part
            for field in row
            if field not in core_fields and field not in INTERNAL_FIELDS
        }
    )
    columns = EXPORT_COLUMNS + [(field, field.replace("_", " ").title()) for field in extra_fields]
    paths: list[Path] = []
    for index, rows in enumerate(parts, 1):
        workbook = _styled_workbook(rows, columns, sheet_name="Resultados", table_name="TabelaResultados")
        headers = [label for _, label in columns]
        sheet = workbook["Resultados"]
        if "Número" in headers:
            bib_column = headers.index("Número") + 1
            for row_number in range(2, sheet.max_row + 1):
                cell = sheet.cell(row_number, bib_column)
                if cell.value is not None:
                    cell.value = str(cell.value)
                cell.number_format = "@"
        suffix = f"_parte_{index:03d}" if len(parts) > 1 else ""
        path = output_dir / f"openresults_resultados{suffix}.xlsx"
        workbook.save(path)
        paths.append(path)
    if len(paths) == 1:
        return paths[0], paths[0].name
    archive = output_dir / "openresults_resultados.zip"
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in paths:
            bundle.write(path, arcname=path.name)
    for path in paths:
        path.unlink(missing_ok=True)
    return archive, archive.name


INTERNAL_FIELDS = {"category_code", "modality_value"}


def safe_slug(value: str, fallback: str = "resultados") -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", ascii_value).strip("-").lower()
    return (slug or fallback)[:120]


def _excel_value(value: Any) -> Any:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value


def export_xlsx(result: ExtractionResult, output_dir: Path) -> tuple[Path, str]:
    if not result.records:
        raise ExportError("Não há resultados para gerar o arquivo Excel.")
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
        filename = f"{safe_slug(result.metadata.slug or result.metadata.name)}-resultados.xlsx"
        output_path = output_dir / filename

        frame = pd.DataFrame(result.records)
        core_fields = [field for field, _ in EXPORT_COLUMNS]
        extra_fields = sorted(
            field
            for field in frame.columns
            if field not in core_fields and field not in INTERNAL_FIELDS
        )
        ordered_fields = core_fields + extra_fields
        for field in ordered_fields:
            if field not in frame.columns:
                frame[field] = None
        frame = frame[ordered_fields]
        display_names = dict(EXPORT_COLUMNS)
        display_names.update({field: field.replace("_", " ").title() for field in extra_fields})

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Resultados"
        headers = [display_names[field] for field in ordered_fields]
        sheet.append(headers)

        for record in frame.to_dict(orient="records"):
            values: list[Any] = []
            for field in ordered_fields:
                value = _excel_value(record.get(field))
                if field == "bib" and value is not None:
                    value = str(value)
                elif field == "event_date" and isinstance(value, str) and value:
                    try:
                        value = date.fromisoformat(value)
                    except ValueError:
                        pass
                elif field == "extracted_at" and isinstance(value, str) and value:
                    try:
                        value = datetime.fromisoformat(value).replace(tzinfo=None)
                    except ValueError:
                        pass
                values.append(value)
            sheet.append(values)

        last_column = get_column_letter(len(headers))
        last_row = sheet.max_row
        table_ref = f"A1:{last_column}{last_row}"
        table = Table(displayName="TabelaResultados", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
        sheet.auto_filter.ref = table_ref
        sheet.freeze_panes = "A2"

        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill

        bib_column = headers.index("Número") + 1
        date_column = headers.index("Data do evento") + 1
        extracted_column = headers.index("Data e hora da extração") + 1
        for row in range(2, last_row + 1):
            sheet.cell(row, bib_column).number_format = "@"
            sheet.cell(row, date_column).number_format = "dd/mm/yyyy"
            sheet.cell(row, extracted_column).number_format = "dd/mm/yyyy hh:mm:ss"

        for index, header in enumerate(headers, start=1):
            max_length = len(header)
            for row in range(2, min(last_row, 500) + 1):
                value = sheet.cell(row, index).value
                max_length = max(max_length, len(str(value)) if value is not None else 0)
            sheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 10), 45)

        workbook.save(output_path)
        return output_path, filename
    except ExportError:
        raise
    except Exception as exc:
        raise ExportError("Não foi possível gerar o arquivo Excel.") from exc
